"""Import SMS Aero cabinet export files into GLAME customer message history."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple, TYPE_CHECKING

if TYPE_CHECKING:
    from sqlalchemy.ext.asyncio import AsyncSession

SMS_AERO_IMPORT_NAMESPACE = uuid.UUID("74e6d9ab-0228-4c2c-82c2-2af8cc1e7ed1")

STATUS_LABELS = {
    0: "В очереди",
    1: "Доставлено",
    2: "Не доставлено",
    3: "Передано оператору",
    6: "Отклонено",
    8: "На модерации",
}
STATUS_TO_MESSAGE = {
    0: "queued",
    1: "delivered",
    2: "failed",
    3: "sent",
    6: "failed",
    8: "moderation",
}


@dataclass(frozen=True)
class SmsAeroImportRow:
    source_row_number: int
    phone_raw: str
    phone_normalized: str
    text: str
    status: str
    status_code: Optional[int]
    status_label: Optional[str]
    extend_status: Optional[str]
    provider_id: Optional[str]
    sent_at: Optional[datetime]
    raw: Dict[str, Any]
    is_technical: bool = False

    @property
    def external_key(self) -> str:
        if self.provider_id:
            return f"sms_aero:{self.provider_id}"
        stable = "|".join([
            self.phone_normalized,
            self.text,
            self.status,
        ])
        digest = hashlib.sha256(stable.encode("utf-8")).hexdigest()[:24]
        return f"sms_aero:file:{digest}"


def normalize_sms_phone(phone: Optional[Any]) -> Optional[str]:
    if phone is None:
        return None
    clean_phone = "".join(c for c in str(phone) if c.isdigit())
    if len(clean_phone) == 11 and clean_phone.startswith("8"):
        clean_phone = "7" + clean_phone[1:]
    elif len(clean_phone) == 10:
        clean_phone = "7" + clean_phone
    return clean_phone if len(clean_phone) == 11 else None


def is_technical_sms_text(text: str) -> bool:
    value = (text or "").strip().lower()
    if not value:
        return True
    return (
        "код входа glame" in value
        or value == "проверка связи"
        or value.startswith("glame: podarochnyj sertifikat")
        or "pin" in value and "sertifikat" in value
    )


def _canonical_header(value: Any) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    text = re.sub(r"[^a-zа-я0-9]+", "_", text)
    return text.strip("_")


def _first_value(row: Dict[str, Any], aliases: Iterable[str]) -> Optional[Any]:
    canonical = {_canonical_header(k): v for k, v in row.items()}
    for alias in aliases:
        key = _canonical_header(alias)
        if key in canonical and canonical[key] not in (None, ""):
            return canonical[key]
    return None


def _map_status(value: Any, extend_status: Any = None) -> Tuple[str, Optional[int], Optional[str]]:
    raw = str(value if value is not None else "").strip().lower().replace("ё", "е")
    ext = str(extend_status if extend_status is not None else "").strip().lower()
    if raw:
        try:
            code = int(float(raw))
            return STATUS_TO_MESSAGE.get(code, "sent"), code, STATUS_LABELS.get(code)
        except ValueError:
            pass
    joined = f"{raw} {ext}"
    if any(token in joined for token in ["не достав", "undeliver", "failed", "ошиб", "отклон", "reject"]):
        return "failed", 2, STATUS_LABELS[2]
    if any(token in joined for token in ["достав", "delivery", "delivered"]):
        return "delivered", 1, STATUS_LABELS[1]
    if any(token in joined for token in ["очеред", "queue"]):
        return "queued", 0, STATUS_LABELS[0]
    if any(token in joined for token in ["модерац", "moderation"]):
        return "moderation", 8, STATUS_LABELS[8]
    if any(token in joined for token in ["передано", "sent"]):
        return "sent", 3, STATUS_LABELS[3]
    return "sent", None, None


def _parse_sent_at(value: Any) -> Optional[datetime]:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    text = str(value).strip()
    if not text:
        return None
    if text.isdigit() and len(text) >= 9:
        try:
            return datetime.fromtimestamp(int(text), timezone.utc)
        except Exception:
            pass
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


def _parse_sent_at_from_status(*values: Any) -> Optional[datetime]:
    text = " ".join(str(value or "") for value in values)
    match = re.search(r"(\d{2}\.\d{2}\.\d{4})\s*(?:в|,)?\s*(\d{1,2}:\d{2}(?::\d{2})?)", text, flags=re.IGNORECASE)
    if not match:
        return None
    return _parse_sent_at(f"{match.group(1)} {match.group(2)}")


def _rows_from_csv(content: bytes) -> List[Dict[str, Any]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
    if rows and len(rows[0].keys()) == 1:
        header = next(iter(rows[0].keys()))
        if "," in header:
            rows = list(csv.DictReader(io.StringIO(text), delimiter=","))
        elif ";" in header:
            rows = list(csv.DictReader(io.StringIO(text), delimiter=";"))
    return rows


def _rows_from_xlsx(content: bytes) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = next(rows_iter, None)
    if not headers:
        return []
    header_values = [str(h or "").strip() for h in headers]
    rows: List[Dict[str, Any]] = []
    for values in rows_iter:
        row = {header_values[i]: values[i] if i < len(values) else None for i in range(len(header_values))}
        rows.append(row)
    return rows


def parse_sms_aero_import_file(filename: str, content: bytes) -> List[SmsAeroImportRow]:
    suffix = (filename or "").lower().rsplit(".", 1)[-1]
    if suffix in {"xlsx", "xlsm"}:
        raw_rows = _rows_from_xlsx(content)
    elif suffix in {"csv", "txt"}:
        raw_rows = _rows_from_csv(content)
    else:
        raise ValueError("Поддерживаются файлы .xlsx, .csv или .txt из экспорта SMS Aero")

    parsed: List[SmsAeroImportRow] = []
    for idx, row in enumerate(raw_rows, start=2):
        phone_raw = _first_value(row, ["Телефон", "Номер телефона", "Номер", "number", "phone", "recipient", "Кому"])
        phone_normalized = normalize_sms_phone(phone_raw)
        text = str(_first_value(row, ["Текст", "Текст сообщения", "Сообщение", "text", "message", "SMS"]) or "").strip()
        if not phone_normalized or not text:
            continue
        extend_status = _first_value(row, ["Расширенный статус", "extendStatus", "extend_status", "Детальный статус"])
        status_raw = _first_value(row, ["Статус", "status", "status_code", "Код статуса"])
        status, status_code, status_label = _map_status(status_raw, extend_status)
        provider_id = _first_value(row, ["ID", "id", "sms_id", "SMS ID", "Идентификатор"])
        sent_at = _parse_sent_at(_first_value(row, ["Дата отправки", "Дата создания", "Дата и время", "dateSend", "date_send", "Отправлено", "Дата", "dateCreate"]))
        if sent_at is None:
            sent_at = _parse_sent_at_from_status(status_raw, extend_status)
        parsed.append(
            SmsAeroImportRow(
                source_row_number=idx,
                phone_raw=str(phone_raw or ""),
                phone_normalized=phone_normalized,
                text=text,
                status=status,
                status_code=status_code,
                status_label=status_label,
                extend_status=str(extend_status) if extend_status not in (None, "") else None,
                provider_id=str(provider_id) if provider_id not in (None, "") else None,
                sent_at=sent_at,
                raw={str(k): v for k, v in row.items()},
                is_technical=is_technical_sms_text(text),
            )
        )
    return parsed


def _message_id_for_import(row: SmsAeroImportRow) -> uuid.UUID:
    return uuid.uuid5(SMS_AERO_IMPORT_NAMESPACE, row.external_key)


def _generated_messages_base_dir() -> Path:
    project_root = Path(__file__).parent.parent.parent
    base_dir = project_root / "backend" / "generated_messages"
    base_dir.mkdir(parents=True, exist_ok=True)
    return base_dir


def _write_sms_aero_result_file(
    *,
    generation_id: str,
    campaign_id: uuid.UUID,
    campaign_name: str,
    filename: str,
    event_store: Optional[str],
    messages: List[Dict[str, Any]],
    parsed_count: int,
    matched_count: int,
    unmatched_count: int,
    technical_count: int,
    status_counts: Dict[str, int],
) -> str:
    payload = {
        "status": "success",
        "source": "sms_aero_file_import",
        "generation_id": generation_id,
        "campaign_id": str(campaign_id),
        "campaign_name": campaign_name,
        "import_filename": filename,
        "event_store": event_store,
        "count": len(messages),
        "parsed": parsed_count,
        "matched": matched_count,
        "unmatched": unmatched_count,
        "technical_count": technical_count,
        "status_counts": status_counts,
        "messages": messages,
    }
    filepath = _generated_messages_base_dir() / f"messages_{generation_id}.json"
    filepath.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return str(filepath)


async def import_sms_aero_file_to_customer_history(
    db: "AsyncSession",
    *,
    filename: str,
    content: bytes,
    campaign_name: Optional[str] = None,
    event_store: Optional[str] = None,
    event_type: str = "external_sms_aero",
) -> Dict[str, Any]:
    from sqlalchemy import select
    from sqlalchemy.dialects.postgresql import insert as pg_insert

    from app.models.customer_message import CustomerMessage
    from app.models.marketing_campaign import MarketingCampaign
    from app.models.user import User
    from app.services.generation_history import get_generation_history

    rows = parse_sms_aero_import_file(filename, content)
    # Сохраняем все строки импорта, включая технические SMS (OTP/сертификаты).
    # Технические сообщения помечаются в payload и могут фильтроваться в UI/аналитике,
    # но больше не теряются при загрузке файла.
    import_rows = rows
    effective_campaign_name = (campaign_name or filename.rsplit(".", 1)[0] or "SMS Aero импорт").strip()
    campaign_seed = f"{filename}|{effective_campaign_name}|{import_rows[0].text if import_rows else ''}"
    campaign_id = uuid.uuid5(SMS_AERO_IMPORT_NAMESPACE, f"campaign:{campaign_seed}")
    generation_id = str(campaign_id)
    sent_dates = [row.sent_at for row in import_rows if row.sent_at]
    campaign_start = min(sent_dates) if sent_dates else datetime.now(timezone.utc)
    campaign_end = max(sent_dates) if sent_dates else campaign_start
    phone_set = {row.phone_normalized for row in import_rows}

    users_by_phone: Dict[str, User] = {}
    if phone_set:
        result = await db.execute(select(User).where(User.is_customer == True, User.phone.isnot(None)))
        for user in result.scalars().all():
            normalized = normalize_sms_phone(getattr(user, "phone", None))
            if normalized in phone_set and normalized not in users_by_phone:
                users_by_phone[normalized] = user

    now = datetime.now(timezone.utc)
    insert_rows: List[Dict[str, Any]] = []
    result_messages: List[Dict[str, Any]] = []
    unmatched = 0
    status_counts: Dict[str, int] = {}
    technical_count = 0
    for row in import_rows:
        if row.is_technical:
            technical_count += 1
        status_counts[row.status] = status_counts.get(row.status, 0) + 1
        user = users_by_phone.get(row.phone_normalized)
        if not user:
            unmatched += 1
            continue
        delivery = {
            "provider": "sms_aero",
            "sms_id": row.provider_id,
            "status": row.status,
            "status_code": row.status_code,
            "status_label": row.status_label,
            "extend_status": row.extend_status,
            "last_checked_at": now.isoformat(),
            "last_response": {
                "id": row.provider_id,
                "status": row.status_code,
                "extendStatus": row.extend_status,
                "dateSend": row.sent_at.isoformat() if row.sent_at else None,
            },
        }
        payload = {
            "source": "sms_aero_file_import",
            "message_kind": "broadcast",
            "generation_id": generation_id,
            "campaign_id": str(campaign_id),
            "campaign_name": effective_campaign_name,
            "event_store": event_store,
            "phone": row.phone_normalized,
            "sms_id": row.provider_id,
            "sms_delivery": delivery,
            "sms_status_checked_at": now.isoformat(),
            "sms_aero_external_key": row.external_key,
            "source_row_number": row.source_row_number,
            "import_filename": filename,
            "is_technical": row.is_technical,
            "raw": row.raw,
        }
        result_messages.append({
            "client_id": str(user.id),
            "user_id": str(user.id),
            "phone": row.phone_normalized,
            "name": getattr(user, "full_name", None),
            "gender": getattr(user, "gender", None),
            "segment": None,
            "event_type": event_type,
            "event_brand": None,
            "event_store": event_store,
            "store": event_store,
            "message": row.text,
            "cta": None,
            "status": row.status,
            "sent_at": (row.sent_at or campaign_start).isoformat(),
            "campaign_id": str(campaign_id),
            "campaign_name": effective_campaign_name,
            "sms_id": row.provider_id,
            "sms_status": row.status,
            "sms_status_code": row.status_code,
            "sms_status_label": row.status_label,
            "sms_extend_status": row.extend_status,
            "source_row_number": row.source_row_number,
            "is_technical": row.is_technical,
            "raw": row.raw,
        })
        insert_rows.append({
            "id": _message_id_for_import(row),
            "user_id": user.id,
            "message": row.text,
            "cta": None,
            "segment": None,
            "event_type": event_type,
            "event_brand": None,
            "event_store": event_store,
            "payload": payload,
            "status": row.status,
            "sent_at": row.sent_at or campaign_start,
        })

    if insert_rows:
        stmt = pg_insert(CustomerMessage).values(insert_rows)
        stmt = stmt.on_conflict_do_update(
            index_elements=["id"],
            set_={
                "message": stmt.excluded.message,
                "event_store": stmt.excluded.event_store,
                "payload": stmt.excluded.payload,
                "status": stmt.excluded.status,
                "sent_at": stmt.excluded.sent_at,
            },
        )
        await db.execute(stmt)

    campaign = await db.get(MarketingCampaign, campaign_id)
    campaign_metrics = {
        "source": "sms_aero_file_import",
        "import_filename": filename,
        "recipients": len(insert_rows),
        "delivered": status_counts.get("delivered", 0),
        "failed": status_counts.get("failed", 0),
        "unmatched": unmatched,
        "technical": technical_count,
        "delivery_rate": round(status_counts.get("delivered", 0) / len(import_rows) * 100, 2) if import_rows else 0,
        "message": import_rows[0].text if import_rows else "",
    }
    if campaign is None:
        campaign = MarketingCampaign(
            id=campaign_id,
            name=effective_campaign_name,
            type="sms",
            status="completed",
            start_date=campaign_start,
        )
        db.add(campaign)
    campaign.name = effective_campaign_name
    campaign.type = "sms"
    campaign.status = "completed"
    campaign.start_date = campaign_start
    campaign.end_date = campaign_end
    campaign.channels = ["sms"]
    campaign.target_audience = {"customer_count": len(insert_rows), "store": event_store}
    campaign.metrics = campaign_metrics
    await db.commit()

    saved_file = _write_sms_aero_result_file(
        generation_id=generation_id,
        campaign_id=campaign_id,
        campaign_name=effective_campaign_name,
        filename=filename,
        event_store=event_store,
        messages=result_messages,
        parsed_count=len(rows),
        matched_count=len(insert_rows),
        unmatched_count=unmatched,
        technical_count=technical_count,
        status_counts=status_counts,
    )

    await get_generation_history().upsert_completed(
        generation_id,
        event_type=effective_campaign_name,
        segment=event_store or "SMS Aero",
        started_at=campaign_start.isoformat(),
        completed_at=campaign_end.isoformat(),
        total=len(import_rows),
        success=status_counts.get("delivered", 0),
        errors=status_counts.get("failed", 0) + unmatched,
        params={
            "source": "sms_aero_file_import",
            "campaign_id": str(campaign_id),
            "campaign_name": effective_campaign_name,
            "import_filename": filename,
            "store": event_store,
        },
        saved_file=saved_file,
    )

    return {
        "status": "ok",
        "filename": filename,
        "parsed": len(rows),
        "technical": technical_count,
        "eligible": len(import_rows),
        "matched": len(insert_rows),
        "imported": len(insert_rows),
        "unmatched": unmatched,
        "status_counts": status_counts,
        "campaign_id": str(campaign_id),
        "generation_id": generation_id,
        "campaign_name": effective_campaign_name,
        "saved_file": saved_file,
    }
